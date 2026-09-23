"""The holdings grid: read a book into an editable table, and write it back.

A portfolio is not a document, it is a position that changes. Uploading a fresh
spreadsheet after every trade is how a monitoring tool stops being used by the
second week, so the book has to be editable in place.

What is edited is the position sheet -- one row per holding, carrying the
quantity, the average cost and the date the position was opened -- because that
is the shape every custodian statement arrives in and the shape the analysis
already reads. Trades are applied to it rather than stored beside it: a buy
re-weights the average cost, a sale reduces the quantity, and a sale of the
whole position removes the row.

That last one is a real limitation, and it belongs on screen rather than buried
here: this sheet records what is *held*. A position sold in full leaves the
book, and the record of what it earned leaves with it. Every figure in the tool
is therefore the return on what the client still owns.

Edits are never written over the uploaded file. They go to a separate workbook
under the app's own data root and the registry points at it, so the document the
client sent stays exactly as they sent it -- which is the only version that
settles an argument about a number.
"""

from __future__ import annotations

import time
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from .portfolio import (
    MARKET_CURRENCY, SECTION_WORDS, _country_from_ticker, _find_header_row,
    _map_columns, _norm, _normalise_country, normalise_ticker, parse_date,
    parse_number,
)
from .sectors import market_from_suffix

#: The sheet the analysis reads, and the only one this module rewrites.
SHEET = "Holdings"

#: Column order, matching the canonical workbook written by ingest.py so an
#: edited book is readable by exactly the same loader as an uploaded one.
COLUMNS: list[str] = ["Security Name", "Ticker", "Ccy", "Quantity", "Avg Cost",
                      "Purchase Date", "Country", "Sector"]

NUMERIC = ("Quantity", "Avg Cost")

#: A purchase date before this is a data-entry slip, not a long-term holding.
_EARLIEST = date(1970, 1, 1)


class EditError(ValueError):
    """The grid cannot be saved as it stands, and the reason is for a reader."""


# -- reading ----------------------------------------------------------------

def read_holdings(path: Path | str, sheet: str = SHEET) -> pd.DataFrame:
    """The holdings sheet as a clean rectangle, whatever shape it arrived in.

    The same tolerant scan the loader uses: find the real header under whatever
    title block sits above it, map the column synonyms, and drop the section
    headings and total rows that are structure rather than positions. What comes
    back is what the analyst sees in the grid, so the two can never disagree.
    """
    raw = pd.read_excel(Path(path), sheet_name=sheet, header=None)
    header_row = _find_header_row(raw)
    cols = _map_columns(raw.iloc[header_row].tolist())

    section_country = ""
    rows: list[dict] = []

    for i in range(header_row + 1, len(raw)):
        row = raw.iloc[i].tolist()
        first = (str(row[cols["name"]]).strip()
                 if "name" in cols and pd.notna(row[cols["name"]]) else "")
        if not first or first.lower() == "nan":
            continue

        populated = sum(1 for c in row if pd.notna(c) and str(c).strip())
        if populated <= 2 and _norm(first) in {_norm(s) for s in SECTION_WORDS}:
            section_country = {"unitedstates": "US", "india": "IN",
                               "saudiarabia": "SA"}.get(_norm(first),
                                                        section_country)
            continue
        if _norm(first) in {"total", "subtotal", "grandtotal"}:
            continue

        country = section_country
        if "country" in cols and pd.notna(row[cols["country"]]):
            stated = str(row[cols["country"]]).strip().upper()
            if stated and stated != "NAN":
                country = _normalise_country(stated)

        ticker = normalise_ticker(row[cols["ticker"]], country) if "ticker" in cols else ""
        sector = ""
        if "sector" in cols and pd.notna(row[cols["sector"]]):
            candidate = str(row[cols["sector"]]).strip()
            if candidate.lower() != "nan":
                sector = candidate

        currency = ""
        if "currency" in cols and pd.notna(row[cols["currency"]]):
            stated = str(row[cols["currency"]]).strip().upper()
            if stated and stated != "NAN":
                currency = stated

        # The suffix, not the old "anything unrecognised is American" default.
        # A blank market is then a visible gap the editor offers to fill, which
        # is a great deal better than a Tadawul code quietly filed as a US
        # listing and measured against the S&P.
        market = country or market_from_suffix(ticker)

        rows.append({
            "Security Name": first,
            "Ticker": ticker,
            "Ccy": currency or MARKET_CURRENCY.get(market, ""),
            "Quantity": parse_number(row[cols["quantity"]]) if "quantity" in cols else None,
            "Avg Cost": parse_number(row[cols["cost"]]) if "cost" in cols else None,
            "Purchase Date": parse_date(row[cols["date"]]) if "date" in cols else None,
            "Country": market,
            "Sector": sector,
        })

    return pd.DataFrame(rows, columns=COLUMNS)


def blank_row(ticker: str = "", name: str = "", country: str = "") -> dict:
    """A new position, with what can be inferred from the ticker filled in.

    Unknown stays unknown. Defaulting an unplaceable symbol to a US listing in
    dollars is how a Tadawul code ends up measured against the S&P, and a blank
    the editor offers to fill is the better failure.
    """
    code = (country or market_from_suffix(ticker)).upper()
    return {
        "Security Name": name or ticker,
        "Ticker": ticker,
        "Ccy": MARKET_CURRENCY.get(code, ""),
        "Quantity": None,
        "Avg Cost": None,
        "Purchase Date": date.today(),
        "Country": code,
        "Sector": "",
    }


# -- validating -------------------------------------------------------------

def tidy(frame: pd.DataFrame) -> pd.DataFrame:
    """Normalise what a person typed before anything is checked or saved."""
    out = frame.copy()
    for column in COLUMNS:
        if column not in out.columns:
            out[column] = None
    out = out[COLUMNS]
    if out.empty:
        # An empty frame has no rows to apply anything to, and a row-wise
        # apply over one returns a DataFrame rather than a mask, which then
        # indexes by column and leaves a frame with no columns at all.
        return out.reset_index(drop=True)

    # A row with nothing in it is the grid's own trailing blank, not a holding.
    out = out[~out.apply(lambda r: all(
        pd.isna(v) or str(v).strip() == "" for v in r), axis=1)]

    out["Security Name"] = out["Security Name"].fillna("").astype(str).str.strip()
    out["Country"] = (out["Country"].fillna("").astype(str).str.strip()
                      .str.upper())
    out["Ccy"] = out["Ccy"].fillna("").astype(str).str.strip().str.upper()
    out["Sector"] = out["Sector"].fillna("").astype(str).str.strip()
    out["Ticker"] = [
        normalise_ticker(t, c) for t, c in zip(out["Ticker"], out["Country"])
    ]
    for column in NUMERIC:
        out[column] = pd.to_numeric(out[column], errors="coerce")
    out["Purchase Date"] = [_as_date(v) for v in out["Purchase Date"]]
    return out.reset_index(drop=True)


def _as_date(value) -> date | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(value)


def validate(frame: pd.DataFrame, today: date | None = None) -> list[str]:
    """Every problem with the grid, phrased for the person who has to fix it.

    Returns all of them rather than the first: a save that reports one error at
    a time turns a five-row correction into five round trips.
    """
    today = today or date.today()
    problems: list[str] = []

    if frame.empty:
        return ["The portfolio has no holdings. Add at least one row."]

    for i, row in frame.iterrows():
        where = f"Row {i + 1}"
        name = str(row["Security Name"]).strip()
        label = f"{where} ({name})" if name else where

        if not name:
            problems.append(f"{where}: give the security a name.")
        if not str(row["Ticker"]).strip():
            problems.append(f"{label}: a ticker is required -- it is how prices "
                            "and the benchmark are found.")

        qty = row["Quantity"]
        if pd.isna(qty) or qty <= 0:
            problems.append(f"{label}: quantity must be greater than zero. "
                            "To close a position, delete the row.")

        cost = row["Avg Cost"]
        if not pd.isna(cost) and cost <= 0:
            problems.append(f"{label}: average cost must be greater than zero, "
                            "or left empty to use the closing price on the "
                            "purchase date.")

        when = row["Purchase Date"]
        if when is None:
            problems.append(f"{label}: a purchase date is required -- the "
                            "return depends on when the money went in.")
        elif when > today:
            problems.append(f"{label}: the purchase date {when} is in the "
                            "future.")
        elif when < _EARLIEST:
            problems.append(f"{label}: the purchase date {when} looks like a "
                            "typing slip.")

        ccy = str(row["Ccy"]).strip()
        if len(ccy) != 3 or not ccy.isalpha():
            problems.append(f"{label}: currency should be a three-letter code "
                            f"such as USD or INR, not {ccy or 'blank'}.")

        country = str(row["Country"]).strip()
        if len(country) != 2 or not country.isalpha():
            problems.append(f"{label}: market should be a two-letter code such "
                            f"as US or IN, not {country or 'blank'}.")

    # One ticker, one row. Two rows sharing a ticker become two holdings with
    # the same identity, and everything keyed on it -- the benchmark decision,
    # the news feed, the detail page -- would resolve to whichever came first.
    tickers = [str(t).strip().upper() for t in frame["Ticker"] if str(t).strip()]
    for ticker in sorted({t for t in tickers if tickers.count(t) > 1}):
        problems.append(
            f"{ticker} appears on {tickers.count(ticker)} rows. Combine them "
            "into one position, or the tool cannot tell them apart."
        )

    return problems


@dataclass
class PriceProblem:
    """A ticker that would not price, and what the feed said about it."""

    ticker: str
    reason: str
    #: True when the feed states there is no such symbol. Anything else -- a
    #: rate limit, a timeout, a DNS failure -- is the plumbing, not the ticker,
    #: and must not be reported as a spelling mistake.
    unknown: bool = False


#: The feed's own words when a symbol does not exist, as opposed to when it
#: could not be reached.
_NO_SUCH_SYMBOL = ("no data returned", "may be delisted", "not found",
                   "no timezone found")


def unpriced(frame: pd.DataFrame, md, only: set[str] | None = None,
             attempts: int = 2, pause: float = 1.5) -> list[PriceProblem]:
    """Tickers that would not price, with the reason for each.

    ``only`` narrows the check to tickers that changed, because the answer for
    one that was already in the book has not.

    Tried more than once, with a wait between attempts: an immediate retry is
    no defence against a rate limiter, which is the failure this actually meets
    in practice. And the reason is carried back rather than swallowed. A guard
    that refuses a trade while hiding why is worse than no guard at all -- it
    sent an analyst to check a ticker that had been right all along.
    """
    problems: list[PriceProblem] = []
    tries = max(attempts, 1)

    for ticker in frame["Ticker"]:
        code = str(ticker).strip()
        if not code or (only is not None and code not in only):
            continue

        last = ""
        for attempt in range(tries):
            try:
                md.prices(code)
                last = ""
                break
            except Exception as exc:  # noqa: BLE001 - the feed raises broadly
                last = str(exc) or type(exc).__name__
                if attempt < tries - 1:
                    time.sleep(pause)

        if last:
            lowered = last.lower()
            problems.append(PriceProblem(
                ticker=code, reason=last,
                unknown=any(p in lowered for p in _NO_SUCH_SYMBOL)))
    return problems


# -- sectors ----------------------------------------------------------------

#: The columns worth chasing when they are blank, and why each one matters.
#: Both decide a benchmark, which is why neither is treated as cosmetic.
GAP_FIELDS: dict[str, str] = {
    "Sector": "sector",
    "Country": "market",
}


def missing(frame: pd.DataFrame, column: str) -> list[tuple[str, str, str]]:
    """Rows with nothing in ``column``, as (ticker, name, market)."""
    out: list[tuple[str, str, str]] = []
    for _, row in tidy(frame).iterrows():
        ticker = str(row["Ticker"]).strip()
        if ticker and not str(row[column]).strip():
            out.append((ticker, str(row["Security Name"]).strip(),
                        str(row["Country"]).strip()))
    return out


def missing_sectors(frame: pd.DataFrame) -> list[tuple[str, str, str]]:
    """Rows with no sector, as (ticker, name, market) ready to look up.

    A blank sector is not a cosmetic gap. It decides the benchmark: with one,
    an Indian bank is measured against the Nifty Bank index; without one it
    falls back to the broad market and its industry is never isolated.
    """
    return missing(frame, "Sector")


def missing_markets(frame: pd.DataFrame) -> list[tuple[str, str, str]]:
    """Rows with no market.

    The market decides more than the sector does. It picks which country's
    indices the holding is even eligible for -- the whole point of measuring an
    Indian bank against Bank Nifty rather than against global financials -- and
    it decides the currency. A blank one is the more expensive of the two.
    """
    return missing(frame, "Country")


def row_gaps(row, misses: dict) -> str:
    """What still needs a look on one row, as a short label, or "" if
    nothing does.

    Ticker/Quantity/Purchase Date mirror validate()'s own required fields --
    a save would refuse a blank one anyway, so it is worth a flag before that
    point, not just at it. Country/Sector are flagged only once something has
    genuinely tried and failed to place them (misses), or the cell is blank
    outright; a value already stated in the file is never called a gap.
    """
    ticker = str(row.get("Ticker", "")).strip()
    parts = []
    if not str(row.get("Security Name", "")).strip():
        parts.append("Security Name")
    if not ticker:
        parts.append("Ticker")
    qty = row.get("Quantity")
    if qty is None or (isinstance(qty, float) and pd.isna(qty)) or qty == 0:
        parts.append("Quantity")
    if row.get("Purchase Date") is None:
        parts.append("Purchase Date")
    if not str(row.get("Country", "")).strip() or (ticker, "Country") in misses:
        parts.append("Country")
    if not str(row.get("Sector", "")).strip() or (ticker, "Sector") in misses:
        parts.append("Sector")
    return ", ".join(parts)


def review_order(frame: pd.DataFrame, misses: dict) -> list[str]:
    """Row order for a just-uploaded book: every ticker with a gap first, in
    the file's own order, then everything else. Frozen at upload time rather
    than recomputed on every keystroke, so fixing the one cell being looked
    at does not also move that row out from under the cursor.
    """
    tickers = [str(t).strip() for t in frame["Ticker"]]
    flagged = [t for t, (_, row) in zip(tickers, frame.iterrows())
               if row_gaps(row, misses)]
    flagged_set = set(flagged)
    rest = [t for t in tickers if t not in flagged_set]
    return flagged + rest


def fill_sectors(frame: pd.DataFrame, lookup):
    """Fill blank sectors using ``lookup(ticker, name, country)``.

    Returns the frame, what was filled, and -- just as important -- what was
    asked for and not found, each with the reason. A holding that no source
    could place has to be distinguishable from one nobody has looked up yet;
    otherwise the only offer the screen can make is to run the same failing
    lookup again.

    Only blanks are touched. A sector the client stated, or an analyst typed,
    is theirs and outranks any feed -- they may know something about the
    business that a classification table does not.

    The lookup is injected rather than imported so that this stays a pure
    frame-to-frame function: the network lives in sectors.py, and the tests
    here never reach for it.
    """
    return _fill(frame, lookup, "Sector", "sector", "attribution")


def fill_markets(frame: pd.DataFrame, lookup):
    """Fill blank markets from the same lookup, same rules.

    Same shape as fill_sectors, and cheap to run beside it: the lookup caches
    per ticker, so asking it for a market costs nothing once the sector has
    been asked for.
    """
    return _fill(frame, lookup, "Country", "market", "market_attribution")


def _fill(frame: pd.DataFrame, lookup, column: str, attribute: str,
          credit: str):
    out = tidy(frame)
    filled: list[tuple[str, str, str]] = []
    unresolved: list[tuple[str, str, str]] = []

    for ticker, name, country in missing(out, column):
        try:
            guess = lookup(ticker, name, country)
        except Exception as exc:  # noqa: BLE001 - a failed source is a reason
            unresolved.append((ticker, name, f"the lookup failed: {exc}"))
            continue
        value = getattr(guess, attribute, "") if guess is not None else ""
        if not value:
            unresolved.append((ticker, name,
                               getattr(guess, "note", "")
                               or "no source could place this company"))
            continue
        out.loc[out["Ticker"].str.upper() == ticker.upper(), column] = value
        filled.append((ticker, value,
                       getattr(guess, credit, "") or "identified"))

    return out, filled, unresolved


def set_values(frame: pd.DataFrame, chosen: dict[str, str],
               column: str) -> pd.DataFrame:
    """Write values chosen by hand, keyed by ticker.

    Separate from the fill functions because it is the opposite act: those ask
    a source and defer to what the book already says, this is the analyst
    overruling every source there is. Blank choices are ignored, so a partly
    completed form sets what was answered and leaves the rest alone.
    """
    out = tidy(frame)
    for ticker, value in chosen.items():
        label = str(value or "").strip()
        code = str(ticker or "").strip().upper()
        if not label or not code:
            continue
        out.loc[out["Ticker"].str.upper() == code, column] = label
    return out


def set_sectors(frame: pd.DataFrame, chosen: dict[str, str]) -> pd.DataFrame:
    return set_values(frame, chosen, "Sector")


def set_markets(frame: pd.DataFrame, chosen: dict[str, str]) -> pd.DataFrame:
    return set_values(frame, chosen, "Country")


def _amount(value) -> str:
    if value is None or pd.isna(value):
        return "blank"
    number = float(value)
    return f"{number:,.0f}" if abs(number - round(number)) < 1e-9 else f"{number:,.2f}"


def changes(before: pd.DataFrame, after: pd.DataFrame) -> list[str]:
    """What editing the grid would do, in plain sentences.

    Shown before the save rather than after. A deleted row is a position and
    its whole history leaving the analysis, and the moment to notice that is
    while it can still be undone by pressing nothing.
    """
    old = {str(r["Ticker"]).upper(): r for _, r in tidy(before).iterrows()}
    new = {str(r["Ticker"]).upper(): r for _, r in tidy(after).iterrows()}
    out: list[str] = []

    for ticker in sorted(new.keys() - old.keys()):
        row = new[ticker]
        out.append(f"<b>{row['Security Name'] or ticker}</b> added: "
                   f"{_amount(row['Quantity'])} shares at "
                   f"{_amount(row['Avg Cost'])}.")

    for ticker in sorted(old.keys() - new.keys()):
        row = old[ticker]
        out.append(f"<b>{row['Security Name'] or ticker}</b> removed. Its "
                   "return leaves every figure in the tool with it.")

    for ticker in sorted(new.keys() & old.keys()):
        was, now = old[ticker], new[ticker]
        label = now["Security Name"] or ticker
        for column, wording in (("Quantity", "quantity"),
                                ("Avg Cost", "average cost")):
            a, b = was[column], now[column]
            if pd.isna(a) and pd.isna(b):
                continue
            if pd.isna(a) or pd.isna(b) or abs(float(a) - float(b)) > 1e-9:
                out.append(f"<b>{label}</b> {wording} {_amount(a)} to "
                           f"{_amount(b)}.")
        if was["Purchase Date"] != now["Purchase Date"]:
            out.append(f"<b>{label}</b> purchase date {was['Purchase Date']} "
                       f"to {now['Purchase Date']}.")
        for column, wording in (("Security Name", "name"), ("Ccy", "currency"),
                                ("Country", "market"), ("Sector", "sector")):
            if str(was[column]).strip() != str(now[column]).strip():
                out.append(f"<b>{label}</b> {wording} "
                           f"{was[column] or 'blank'} to "
                           f"{now[column] or 'blank'}.")

    return out


# -- trades -----------------------------------------------------------------

def apply_trade(frame: pd.DataFrame, ticker: str, side: str, quantity: float,
                price: float, when: date, *, name: str = "",
                country: str = "", sector: str = "",
                currency: str = "") -> tuple[pd.DataFrame, str]:
    """Apply one buy or sale to the position sheet, and say what it did.

    A buy re-weights the average cost across the old and new shares, which is
    the arithmetic people most often get wrong by hand and the reason this is a
    form rather than a note telling the analyst to work it out. The purchase
    date stays at the position's own opening, because that is what every
    "since inception" figure in the tool is measured from.

    A sale leaves the average cost alone -- selling does not change what the
    remaining shares cost -- and removes the row entirely when the last share
    goes, since this sheet holds open positions.
    """
    out = tidy(frame)
    code = normalise_ticker(ticker, country)
    if not code:
        raise EditError("Give a ticker.")
    if quantity is None or quantity <= 0:
        raise EditError("Quantity must be greater than zero.")

    matches = out.index[out["Ticker"].str.upper() == code.upper()].tolist()
    buying = side.lower().startswith("b")

    if not matches:
        if not buying:
            raise EditError(f"{code} is not in this portfolio, so there is "
                            "nothing to sell.")
        if price is None or price <= 0:
            raise EditError("Give the price paid per share.")
        row = blank_row(code, name, country)
        row.update({"Quantity": float(quantity), "Avg Cost": float(price),
                    "Purchase Date": when, "Sector": sector or ""})
        if currency:
            row["Ccy"] = currency.upper()
        out = pd.concat([out, pd.DataFrame([row])], ignore_index=True)
        return tidy(out), (f"Opened {code}: {quantity:,.0f} shares at "
                           f"{price:,.2f} on {when}.")

    at = matches[0]
    held = float(out.at[at, "Quantity"] or 0.0)
    cost = out.at[at, "Avg Cost"]
    cost = float(cost) if not pd.isna(cost) else None
    label = str(out.at[at, "Security Name"]) or code

    if buying:
        if price is None or price <= 0:
            raise EditError("Give the price paid per share.")
        total = held + float(quantity)
        if cost is None:
            blended = float(price)
        else:
            blended = (held * cost + float(quantity) * float(price)) / total
        out.at[at, "Quantity"] = total
        out.at[at, "Avg Cost"] = blended
        return tidy(out), (
            f"Bought {quantity:,.0f} more {label} at {price:,.2f}. "
            f"Position {held:,.0f} to {total:,.0f} shares, average cost "
            f"{cost:,.2f} to {blended:,.2f}." if cost is not None else
            f"Bought {quantity:,.0f} more {label} at {price:,.2f}. "
            f"Position {held:,.0f} to {total:,.0f} shares.")

    if float(quantity) > held + 1e-9:
        raise EditError(f"{label} holds {held:,.0f} shares; you cannot sell "
                        f"{float(quantity):,.0f}.")

    remaining = held - float(quantity)
    if remaining <= 1e-9:
        out = out.drop(index=at)
        return tidy(out), (
            f"Sold the whole {label} position ({held:,.0f} shares). The row is "
            "gone, and with it the record of what it earned -- this sheet "
            "holds what is still owned.")

    out.at[at, "Quantity"] = remaining
    return tidy(out), (f"Sold {quantity:,.0f} {label}. Position {held:,.0f} to "
                       f"{remaining:,.0f} shares; average cost unchanged at "
                       f"{cost:,.2f}." if cost is not None else
                       f"Sold {quantity:,.0f} {label}. Position {held:,.0f} to "
                       f"{remaining:,.0f} shares.")


# -- writing ----------------------------------------------------------------

def _cell(value):
    """Excel-safe: NaN is not a value, it is the absence of one."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    return value


def write_holdings(frame: pd.DataFrame, out_path: Path | str,
                   source: Path | str | None = None) -> Path:
    """Write the grid to a workbook, keeping every other sheet intact.

    The source book is opened and only its Holdings sheet is replaced, so a
    client file carrying fund commitments or property alongside the equities
    does not lose them the first time someone corrects a share count.
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    clean = tidy(frame)

    book = None
    if source and Path(source).exists():
        try:
            book = load_workbook(Path(source))
        except Exception:  # noqa: BLE001 - an unreadable source starts fresh
            book = None
    if book is None:
        book = Workbook()
        book.remove(book.active)

    index = book.sheetnames.index(SHEET) if SHEET in book.sheetnames else 0
    if SHEET in book.sheetnames:
        del book[SHEET]
    sheet = book.create_sheet(SHEET, index)

    sheet.append(list(COLUMNS))
    for _, row in clean.iterrows():
        sheet.append([_cell(row[column]) for column in COLUMNS])

    book.save(out_path)
    return out_path
