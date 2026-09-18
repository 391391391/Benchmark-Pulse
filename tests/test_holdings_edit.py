"""Tests for editing the holdings sheet in place.

The point of the editor is that a portfolio can be kept current without
re-uploading a spreadsheet after every trade. That only works if three things
hold: whatever shape the file arrived in can be read into a clean grid, the
grid can be written back without destroying the rest of the workbook, and
nothing that would silently corrupt the analysis can be saved. A quantity of
zero, two rows claiming the same ticker or a purchase dated next month all
produce a book that loads and lies, which is worse than one that refuses.
"""

from __future__ import annotations

from datetime import date, timedelta

import pandas as pd
import pytest
from openpyxl import load_workbook

from benchmark_pulse.holdings_edit import (
    COLUMNS, EditError, apply_trade, blank_row, changes, read_holdings, tidy,
    unpriced, validate, write_holdings,
)

TODAY = date(2026, 9, 17)


def frame(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=COLUMNS)


def holding(ticker="AAPL", name="Apple Inc.", qty=1000.0, cost=150.0,
            when=date(2022, 5, 16), ccy="USD", country="US",
            sector="Technology") -> dict:
    return {"Security Name": name, "Ticker": ticker, "Ccy": ccy,
            "Quantity": qty, "Avg Cost": cost, "Purchase Date": when,
            "Country": country, "Sector": sector}


def book(tmp_path, rows: list[dict], *, title_block: bool = False,
         extra_sheet: bool = False):
    """A workbook on disk, optionally as untidy as a real custodian file."""
    path = tmp_path / "book.xlsx"
    holdings = frame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        if title_block:
            # Two junk rows above the header, which is what arrives in practice
            # and what the tolerant reader exists for.
            pad = pd.DataFrame([["Quarterly Statement", None, None, None,
                                 None, None, None, None],
                                [None, None, None, None, None, None, None, None],
                                COLUMNS] + holdings.values.tolist())
            pad.to_excel(xl, sheet_name="Holdings", index=False, header=False)
        else:
            holdings.to_excel(xl, sheet_name="Holdings", index=False)
        if extra_sheet:
            pd.DataFrame([{"Fund": "Cedar Growth III", "Vintage": 2021}]).to_excel(
                xl, sheet_name="Fund Commitments", index=False)
    return path


# -- reading ----------------------------------------------------------------

def test_a_clean_sheet_reads_straight_through(tmp_path):
    path = book(tmp_path, [holding(), holding("INFY.NS", "Infosys Ltd",
                                              ccy="INR", country="IN")])
    grid = read_holdings(path)

    assert list(grid.columns) == COLUMNS
    assert list(grid["Ticker"]) == ["AAPL", "INFY.NS"]
    assert grid.loc[1, "Ccy"] == "INR"
    assert grid.loc[0, "Purchase Date"] == date(2022, 5, 16)


def test_a_title_block_above_the_header_is_found_and_skipped(tmp_path):
    """The grid has to show the same rows the analysis loads. If the reader
    here were stricter than the loader, an analyst would edit a table that did
    not match the portfolio on every other screen."""
    path = book(tmp_path, [holding()], title_block=True)
    grid = read_holdings(path)
    assert len(grid) == 1
    assert grid.loc[0, "Ticker"] == "AAPL"


def test_section_headings_and_totals_are_not_holdings(tmp_path):
    path = tmp_path / "grouped.xlsx"
    rows = [COLUMNS,
            ["United States", None, None, None, None, None, None, None],
            ["Apple Inc.", "AAPL", "USD", 1000, 150, date(2022, 5, 16),
             None, "Technology"],
            ["Total", None, None, 1000, None, None, None, None]]
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        pd.DataFrame(rows).to_excel(xl, sheet_name="Holdings", index=False,
                                    header=False)
    grid = read_holdings(path)
    assert list(grid["Security Name"]) == ["Apple Inc."]
    assert grid.loc[0, "Country"] == "US", "the section heading supplies country"


# -- tidying ----------------------------------------------------------------

def test_the_editors_trailing_blank_row_is_not_a_holding():
    out = tidy(frame([holding(), {c: None for c in COLUMNS}]))
    assert len(out) == 1


def test_what_someone_typed_is_normalised_before_it_is_judged():
    out = tidy(frame([holding(ticker=" aapl ", ccy="usd", country="us")]))
    assert out.loc[0, "Ticker"] == "AAPL"
    assert out.loc[0, "Ccy"] == "USD"
    assert out.loc[0, "Country"] == "US"


def test_an_indian_ticker_gains_its_exchange_suffix():
    """'INFY' returns nothing from the price feed; 'INFY.NS' returns the NSE
    listing. Fixing it on the way in beats an error on the way out."""
    out = tidy(frame([holding(ticker="INFY", country="IN")]))
    assert out.loc[0, "Ticker"] == "INFY.NS"


# -- validating -------------------------------------------------------------

def test_a_sound_grid_has_nothing_to_say():
    assert validate(tidy(frame([holding()])), today=TODAY) == []


def test_an_empty_portfolio_is_refused():
    assert validate(tidy(frame([])), today=TODAY)


@pytest.mark.parametrize("field,value,expected", [
    ("Quantity", 0, "greater than zero"),
    ("Quantity", None, "greater than zero"),
    ("Avg Cost", -5, "greater than zero"),
    ("Ticker", "", "ticker is required"),
    ("Security Name", "", "give the security a name"),
    ("Purchase Date", None, "purchase date is required"),
    ("Ccy", "DOLLARS", "three-letter code"),
    ("Country", "USA", "two-letter code"),
])
def test_each_way_a_row_can_be_wrong_is_named(field, value, expected):
    row = holding()
    row[field] = value
    problems = validate(tidy(frame([row])), today=TODAY)
    assert any(expected.lower() in p.lower() for p in problems), problems


def test_a_purchase_dated_in_the_future_is_refused():
    row = holding(when=TODAY + timedelta(days=1))
    problems = validate(tidy(frame([row])), today=TODAY)
    assert any("future" in p for p in problems)


def test_two_rows_cannot_claim_the_same_ticker():
    """Everything downstream is keyed on the ticker -- the benchmark decision,
    the news feed, the detail page. Two rows sharing one would resolve to
    whichever came first and quietly discard the other."""
    problems = validate(tidy(frame([holding(), holding(qty=500)])), today=TODAY)
    assert any("appears on 2 rows" in p for p in problems)


def test_a_blank_average_cost_is_allowed():
    """The loader falls back to the closing price on the purchase date, which
    is a reasonable answer for a position whose cost nobody recorded."""
    assert validate(tidy(frame([holding(cost=None)])), today=TODAY) == []


def test_a_momentary_feed_failure_does_not_condemn_a_good_ticker():
    """The price feed rate-limits, and a refusal looks exactly like a bad
    symbol. Blocking a legitimate trade while telling the analyst to check a
    ticker that was right all along is the worse mistake."""
    class Flaky:
        def __init__(self):
            self.calls = 0

        def prices(self, ticker, **_kw):
            self.calls += 1
            if self.calls == 1:
                raise RuntimeError("429 Too Many Requests")
            return pd.Series([1.0])

    feed = Flaky()
    assert unpriced(frame([holding()]), feed, pause=0) == []
    assert feed.calls == 2, "it should have tried again"


def test_an_outage_is_not_reported_as_a_spelling_mistake():
    """A rate limit is the plumbing, not the ticker. Told apart so the caller
    can warn and carry on rather than refusing the trade."""
    class RateLimited:
        def prices(self, ticker, **_kw):
            raise RuntimeError("429 Too Many Requests")

    trouble = unpriced(frame([holding()]), RateLimited(), pause=0)
    assert [p.ticker for p in trouble] == ["AAPL"]
    assert not trouble[0].unknown
    assert "429" in trouble[0].reason, "the reason has to survive to the screen"


def test_a_symbol_the_feed_does_not_know_is_told_apart():
    class NoSuchThing:
        def prices(self, ticker, **_kw):
            raise RuntimeError(f"{ticker}: no data returned. Check the symbol")

    trouble = unpriced(frame([holding("WRNG")]), NoSuchThing(), pause=0)
    assert trouble[0].unknown, "this one really is a bad ticker"


def test_an_unpriceable_ticker_is_reported_rather_than_dropped():
    class Market:
        def prices(self, ticker, **_kw):
            if ticker == "NOPE":
                raise RuntimeError("NOPE: no data returned")
            return pd.Series([1.0])

    grid = tidy(frame([holding(), holding("NOPE", "Mystery Ltd")]))
    assert [p.ticker for p in unpriced(grid, Market(), pause=0)] == ["NOPE"]
    assert unpriced(grid, Market(), only={"AAPL"}, pause=0) == []


# -- trades -----------------------------------------------------------------

def test_buying_more_reweights_the_average_cost():
    """1000 at 150 plus 500 at 300 is 1500 at 200. Doing this by hand is the
    step people get wrong, which is why it is a form and not a note."""
    out, said = apply_trade(frame([holding()]), "AAPL", "Buy", 500, 300,
                            date(2026, 9, 1))
    assert out.loc[0, "Quantity"] == pytest.approx(1500)
    assert out.loc[0, "Avg Cost"] == pytest.approx(200.0)
    assert "200.00" in said


def test_buying_more_keeps_the_original_purchase_date():
    """Every since-inception figure is measured from it. Moving it to today
    would shorten the holding period and flatter the annualised return."""
    out, _ = apply_trade(frame([holding()]), "AAPL", "Buy", 500, 300,
                         date(2026, 9, 1))
    assert out.loc[0, "Purchase Date"] == date(2022, 5, 16)


def test_buying_something_new_opens_a_position():
    out, said = apply_trade(frame([holding()]), "2222.SR", "Buy", 100, 30,
                            date(2026, 9, 1), name="Saudi Aramco",
                            country="SA")
    row = out[out["Ticker"] == "2222.SR"].iloc[0]
    assert row["Security Name"] == "Saudi Aramco"
    assert row["Ccy"] == "SAR", "the market decides the currency"
    assert row["Quantity"] == pytest.approx(100)
    assert row["Avg Cost"] == pytest.approx(30)
    assert "Opened" in said


def test_selling_part_leaves_the_average_cost_alone():
    """Selling does not change what the remaining shares cost."""
    out, _ = apply_trade(frame([holding()]), "AAPL", "Sell", 400, 0,
                         date(2026, 9, 1))
    assert out.loc[0, "Quantity"] == pytest.approx(600)
    assert out.loc[0, "Avg Cost"] == pytest.approx(150.0)


def test_selling_the_whole_position_removes_the_row_and_says_so():
    out, said = apply_trade(frame([holding(), holding("JPM", "JPMorgan")]),
                            "AAPL", "Sell", 1000, 0, date(2026, 9, 1))
    assert list(out["Ticker"]) == ["JPM"]
    assert "still owned" in said


def test_selling_more_than_is_held_is_refused():
    with pytest.raises(EditError, match="cannot sell"):
        apply_trade(frame([holding()]), "AAPL", "Sell", 2000, 0,
                    date(2026, 9, 1))


def test_selling_something_not_held_is_refused():
    with pytest.raises(EditError, match="nothing to sell"):
        apply_trade(frame([holding()]), "TSLA", "Sell", 10, 0,
                    date(2026, 9, 1))


def test_a_buy_needs_a_price():
    with pytest.raises(EditError, match="price"):
        apply_trade(frame([holding()]), "AAPL", "Buy", 100, 0,
                    date(2026, 9, 1))


def test_a_trade_needs_a_quantity():
    with pytest.raises(EditError, match="greater than zero"):
        apply_trade(frame([holding()]), "AAPL", "Buy", 0, 100,
                    date(2026, 9, 1))


# -- what changed -----------------------------------------------------------

def test_an_added_position_is_described():
    after = pd.concat([frame([holding()]),
                       pd.DataFrame([holding("JPM", "JPMorgan", qty=200)])],
                      ignore_index=True)
    said = changes(frame([holding()]), after)
    assert any("JPMorgan" in s and "added" in s for s in said)


def test_a_removed_position_says_what_goes_with_it():
    said = changes(frame([holding(), holding("JPM", "JPMorgan")]),
                   frame([holding()]))
    assert any("JPMorgan" in s and "removed" in s for s in said)


def test_a_changed_quantity_is_described_both_ways():
    said = changes(frame([holding()]), frame([holding(qty=1500)]))
    assert any("1,000 to 1,500" in s for s in said)


def test_no_change_says_nothing():
    assert changes(frame([holding()]), frame([holding()])) == []


# -- writing ----------------------------------------------------------------

def test_a_grid_survives_the_round_trip(tmp_path):
    path = book(tmp_path, [holding(), holding("INFY.NS", "Infosys Ltd",
                                              ccy="INR", country="IN")])
    grid = read_holdings(path)
    out = write_holdings(grid, tmp_path / "edited.xlsx", source=path)
    again = read_holdings(out)

    pd.testing.assert_frame_equal(tidy(grid), tidy(again))


def test_other_sheets_are_not_collateral_damage(tmp_path):
    """A client book carrying fund commitments beside the equities must not
    lose them the first time someone corrects a share count."""
    path = book(tmp_path, [holding()], extra_sheet=True)
    out = write_holdings(read_holdings(path), tmp_path / "edited.xlsx",
                         source=path)

    sheets = load_workbook(out).sheetnames
    assert "Fund Commitments" in sheets
    assert "Holdings" in sheets


def test_the_written_sheet_is_a_clean_rectangle(tmp_path):
    """Whatever the upload looked like, what is written back has the header on
    row one -- so the next read is trivial and the file is one a person can
    open in Excel and understand."""
    path = book(tmp_path, [holding()], title_block=True)
    out = write_holdings(read_holdings(path), tmp_path / "edited.xlsx",
                         source=path)

    sheet = load_workbook(out)["Holdings"]
    assert [c.value for c in sheet[1]] == COLUMNS
    assert sheet.max_row == 2


def test_writing_without_a_source_still_produces_a_readable_book(tmp_path):
    out = write_holdings(frame([holding()]), tmp_path / "fresh.xlsx")
    assert read_holdings(out).loc[0, "Ticker"] == "AAPL"


def test_a_new_row_is_prefilled_from_the_ticker():
    row = blank_row("2222.SR")
    assert row["Country"] == "SA"
    assert row["Ccy"] == "SAR"
